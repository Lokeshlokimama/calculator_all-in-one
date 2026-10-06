"""Stateless PDF processing service. Documents live only in per-request temp folders."""
import base64
import binascii
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import math
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pypdf import PdfReader, PdfWriter

MAX_BYTES = 10 * 1024 * 1024
JOBS = threading.BoundedSemaphore(1)
ORIGIN = os.environ.get('ALLOWED_ORIGIN', 'https://calculatorsallinone.com')
TOOLS = {'office-to-pdf', 'repair', 'protect', 'unlock', 'pdfa', 'pdf-to-word', 'pdf-to-excel', 'pdf-to-powerpoint', 'markdown', 'compare', 'redact'}

def run(command, cwd):
    try:
        result = subprocess.run(command, cwd=cwd, timeout=90, capture_output=True, check=False)
    except (FileNotFoundError, subprocess.TimeoutExpired) as error:
        raise ValueError('Processing service unavailable or time limit exceeded.') from error
    if result.returncode:
        raise ValueError('This document could not be processed. Check its format and password.')

def decode_file(item, folder, stem):
    name = Path(str(item.get('name', 'document.pdf'))).name
    suffix = Path(name).suffix.lower()
    encoded = item.get('data', '')
    if not isinstance(encoded, str) or len(encoded) > (MAX_BYTES * 4 // 3 + 8):
        raise ValueError('File limit: 10 MB.')
    try:
        data = base64.b64decode(encoded, validate=True)
    except (ValueError, binascii.Error) as error:
        raise ValueError('Invalid file data.') from error
    if not data or len(data) > MAX_BYTES:
        raise ValueError('Choose a non-empty file smaller than 10 MB.')
    if suffix not in {'.pdf', '.doc', '.docx', '.xls', '.xlsx', '.ppt', '.pptx', '.odt', '.ods', '.odp', '.rtf'}:
        raise ValueError('Unsupported file type.')
    source = folder / (stem + suffix)
    source.write_bytes(data)
    return source

def reader(source, password=''):
    if source.suffix != '.pdf' or not source.read_bytes().startswith(b'%PDF-'):
        raise ValueError('Choose a valid PDF.')
    try:
        doc = PdfReader(source, strict=False)
        if doc.is_encrypted and not doc.decrypt(password):
            raise ValueError('A correct password is required. This service does not bypass protection.')
        if not 1 <= len(doc.pages) <= 100:
            raise ValueError('Page limit: 100.')
        return doc
    except ValueError:
        raise
    except Exception as error:
        raise ValueError('Cannot read this PDF.') from error

def process(payload, folder):
    task = payload.get('tool')
    if task not in TOOLS:
        raise ValueError('Unsupported task.')
    source = decode_file(payload.get('file', {}), folder, 'source')
    options = payload.get('options') or {}
    if not isinstance(options, dict):
        raise ValueError('Invalid options.')
    output = folder / 'result.pdf'
    password = str(options.get('password', ''))
    if len(password) > 256:
        raise ValueError('Password is too long.')
    if task == 'office-to-pdf':
        if source.suffix == '.pdf':
            raise ValueError('Choose an Office document.')
        profile = (folder / 'office-profile').as_uri()
        run(['libreoffice', '-env:UserInstallation=' + profile, '--headless', '--convert-to', 'pdf', '--outdir', str(folder), str(source)], folder)
        generated = source.with_suffix('.pdf')
        if not generated.exists():
            raise ValueError('Office conversion failed. The file may be encrypted or unsupported.')
        generated.rename(output)
    elif task == 'repair':
        if source.suffix != '.pdf':
            raise ValueError('Choose a PDF.')
        run(['qpdf', str(source), str(output)], folder)
        reader(output)
    elif task in {'protect', 'unlock'}:
        doc = reader(source, password)
        writer = PdfWriter()
        writer.clone_document_from_reader(doc)
        if task == 'protect':
            new_password = str(options.get('newPassword', ''))
            if not 8 <= len(new_password) <= 256:
                raise ValueError('Use a password containing 8–256 characters.')
            writer.encrypt(new_password, algorithm='AES-256')
        writer.write(output)
    elif task == 'pdfa':
        reader(source, password)
        if password:
            raise ValueError('Unlock the PDF before PDF/A conversion.')
        profiles = list(Path('/usr/share/color/icc').glob('**/srgb.icc')) + list(Path('/usr/share/ghostscript').glob('**/srgb.icc'))
        if not profiles:
            raise ValueError('The PDF/A color profile is not installed.')
        profile = str(profiles[0])
        definition = folder / 'pdfa.ps'
        definition.write_text('/ICCProfile (' + profile + ') def\n[/_objdef {icc_PDFA} /type /stream /OBJ pdfmark\n[{icc_PDFA} << /N 3 >> /PUT pdfmark\n[{icc_PDFA} ICCProfile (r) file /PUT pdfmark\n[/_objdef {OutputIntent_PDFA} /type /dict /OBJ pdfmark\n[{OutputIntent_PDFA} << /Type /OutputIntent /S /GTS_PDFA1 /DestOutputProfile {icc_PDFA} /OutputConditionIdentifier (sRGB) >> /PUT pdfmark\n[{Catalog} << /OutputIntents [ {OutputIntent_PDFA} ] >> /PUT pdfmark\n')
        run(['gs', '-dBATCH', '-dNOPAUSE', '-dSAFER', '--permit-file-read=' + profile, '-sDEVICE=pdfwrite', '-dPDFA=2', '-dPDFACompatibilityPolicy=2', '-sColorConversionStrategy=RGB', '-sOutputFile=' + str(output), str(definition), str(source)], folder)
    elif task in {'markdown', 'compare', 'pdf-to-word', 'pdf-to-excel', 'pdf-to-powerpoint'}:
        doc = reader(source, password)
        pages = [(page.extract_text() or '') for page in doc.pages]
        if sum(map(len, pages)) > 1000000:
            raise ValueError('Text extraction limit exceeded.')
        if task == 'markdown':
            return '\n\n'.join('## Page ' + str(i + 1) + '\n\n' + text for i, text in enumerate(pages)).encode(), 'text/markdown', 'document.md'
        if task == 'compare':
            second = decode_file(payload.get('second', {}), folder, 'second')
            other = reader(second, str(options.get('secondPassword', '')))
            other_pages = [(page.extract_text() or '') for page in other.pages]
            import difflib
            difference = '\n'.join(difflib.unified_diff('\n'.join(pages).splitlines(), '\n'.join(other_pages).splitlines(), fromfile='First PDF', tofile='Second PDF', lineterm=''))
            return (difference or 'No differences in extracted text. Graphics, formatting and images were not compared.').encode(), 'text/plain', 'comparison.txt'
        if task == 'pdf-to-word':
            from docx import Document
            document = Document()
            for index, text in enumerate(pages):
                if index:
                    document.add_page_break()
                for line in text.splitlines():
                    document.add_paragraph(line)
            output = folder / 'result.docx'
            document.save(output)
        elif task == 'pdf-to-excel':
            from openpyxl import Workbook
            book = Workbook()
            book.remove(book.active)
            for index, text in enumerate(pages):
                sheet = book.create_sheet('Page ' + str(index + 1))
                for row, line in enumerate(text.splitlines(), 1):
                    cell = sheet.cell(row, 1)
                    cell.value = line
                    cell.data_type = 's'
            output = folder / 'result.xlsx'
            book.save(output)
        else:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            deck = Presentation()
            for index, text in enumerate(pages):
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                box = slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(9), Inches(6.5))
                box.text_frame.word_wrap = True
                box.text_frame.text = text[:12000]
                if len(text) > 12000:
                    raise ValueError('A page contains too much text for a slide. Use Word export instead.')
                for paragraph in box.text_frame.paragraphs:
                    paragraph.font.size = Pt(14)
            output = folder / 'result.pptx'
            deck.save(output)
    elif task == 'redact':
        doc = reader(source, password)
        if password:
            raise ValueError('Unlock this PDF before redaction.')
        if len(doc.pages) > 20:
            raise ValueError('Raster redaction limit: 20 pages.')
        box = options.get('box', [])
        if len(box) != 4 or any(isinstance(n, bool) or not isinstance(n, (int, float)) or not math.isfinite(n) for n in box):
            raise ValueError('Provide a rectangle as x, y, width, height percentages.')
        x, y, width, height = box
        if not (0 <= x < 100 and 0 <= y < 100 and width > 0 and height > 0 and x + width <= 100 and y + height <= 100):
            raise ValueError('Redaction rectangle is outside the page.')
        run(['pdftoppm', '-scale-to', '1800', '-png', str(source), str(folder / 'page')], folder)
        from PIL import Image, ImageDraw
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import ImageReader
        pdf = canvas.Canvas(str(output))
        images = sorted(folder.glob('page-*.png'), key=lambda p: int(p.stem.split('-')[-1]))
        if len(images) != len(doc.pages):
            raise ValueError('Some pages could not be rendered. No partial file was created.')
        for index, image_path in enumerate(images):
            with Image.open(image_path) as image:
                image = image.convert('RGB')
                w, h = image.size
                ImageDraw.Draw(image).rectangle((x*w/100, y*h/100, (x+width)*w/100, (y+height)*h/100), fill='black')
                page = doc.pages[index]
                pw, ph = float(page.mediabox.width), float(page.mediabox.height)
                pdf.setPageSize((pw, ph))
                pdf.drawImage(ImageReader(image), 0, 0, width=pw, height=ph)
                pdf.showPage()
        pdf.save()
    if not output.exists() or output.stat().st_size > 40 * 1024 * 1024:
        raise ValueError('Output missing or exceeds 40 MB.')
    return output.read_bytes(), 'application/pdf' if output.suffix == '.pdf' else 'application/octet-stream', 'result' + output.suffix

class Handler(BaseHTTPRequestHandler):
    def setup(self):
        super().setup()
        self.connection.settimeout(120)
    def log_message(self, *args):
        pass  # Never log document names, passwords or body data.

    def reply(self, code, body, mime='application/json', filename=None):
        self.send_response(code)
        self.send_header('Content-Type', mime)
        self.send_header('Content-Length', str(len(body)))
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.send_header('Access-Control-Allow-Origin', ORIGIN)
        self.send_header('Access-Control-Expose-Headers', 'Content-Disposition')
        if filename:
            self.send_header('Content-Disposition', 'attachment; filename="' + filename + '"')
        self.end_headers()
        self.wfile.write(body)

    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header('Access-Control-Allow-Origin', ORIGIN)
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()

    def do_GET(self):
        if self.path != '/health':
            return self.reply(404, b'{"error":"Not found"}')
        available = sorted(TOOLS - {'office-to-pdf', 'repair', 'pdfa', 'redact'})
        for task, binary in [('office-to-pdf','libreoffice'), ('repair','qpdf'), ('pdfa','gs'), ('redact','pdftoppm')]:
            if shutil.which(binary):
                available.append(task)
        self.reply(200, json.dumps({'tools':available}).encode())

    def do_POST(self):
        if self.path != '/process':
            return self.reply(404, b'{"error":"Not found"}')
        if self.headers.get('Origin') not in {None, ORIGIN}:
            return self.reply(403, b'{"error":"Origin not allowed"}')
        if not JOBS.acquire(blocking=False):
            return self.reply(429, b'{"error":"Service busy. Try again shortly."}')
        try:
            length = int(self.headers.get('Content-Length', '0'))
            if not 0 < length <= 29 * 1024 * 1024:
                raise ValueError('Request size limit exceeded.')
            if self.headers.get('Content-Type','').split(';')[0] != 'application/json':
                raise ValueError('Use JSON file data.')
            payload = json.loads(self.rfile.read(length))
            if not isinstance(payload, dict):
                raise ValueError('Invalid request.')
            with tempfile.TemporaryDirectory(prefix='pdf-job-') as directory:
                data, mime, filename = process(payload, Path(directory))
                self.reply(200, data, mime, filename)
        except (ValueError, TypeError, KeyError) as error:
            self.reply(400, json.dumps({'error':str(error)}).encode())
        except Exception:
            self.reply(500, b'{"error":"Processing failed. No document was retained."}')
        finally:
            JOBS.release()

if __name__ == '__main__':
    server = ThreadingHTTPServer(('0.0.0.0', int(os.environ.get('PORT','8080'))), Handler)
    server.socket.settimeout(120)
    server.serve_forever()
