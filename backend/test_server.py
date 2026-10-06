import base64
import io
import tempfile
import unittest
import shutil
from pathlib import Path
from server import process
from pypdf import PdfReader
from reportlab.pdfgen import canvas

def fixture(text='First document'):
    stream=io.BytesIO()
    pdf=canvas.Canvas(stream)
    pdf.drawString(50,750,text)
    pdf.showPage()
    pdf.save()
    return {'name':'sample.pdf','data':base64.b64encode(stream.getvalue()).decode()}

class ProcessingTests(unittest.TestCase):
    def call(self,tool,file=None,options=None,second=None):
        with tempfile.TemporaryDirectory() as folder:
            return process({'tool':tool,'file':file or fixture(),'options':options or {},'second':second},Path(folder))
    def test_encryption_requires_correct_password_and_unlocks(self):
        encrypted,_,_=self.call('protect',options={'newPassword':'test-password-123'})
        item={'name':'locked.pdf','data':base64.b64encode(encrypted).decode()}
        with self.assertRaisesRegex(ValueError,'correct password'):
            self.call('unlock',item,{'password':'wrong'})
        result,_,_=self.call('unlock',item,{'password':'test-password-123'})
        doc=PdfReader(io.BytesIO(result))
        self.assertFalse(doc.is_encrypted)
        self.assertIn('First document',doc.pages[0].extract_text())
    def test_markdown_and_compare_text(self):
        result,_,_=self.call('markdown')
        self.assertIn(b'## Page 1',result)
        result,_,_=self.call('compare',second=fixture('Changed document'))
        self.assertIn(b'-First document',result)
        self.assertIn(b'+Changed document',result)
    def test_office_exports_are_readable_documents(self):
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        data,_,_=self.call('pdf-to-word')
        self.assertIn('First document',' '.join(p.text for p in Document(io.BytesIO(data)).paragraphs))
        data,_,_=self.call('pdf-to-excel')
        self.assertEqual(load_workbook(io.BytesIO(data)).active['A1'].value,'First document')
        data,_,_=self.call('pdf-to-powerpoint')
        self.assertIn('First document',Presentation(io.BytesIO(data)).slides[0].shapes[0].text)
    def test_rejects_unsupported_empty_corrupt_and_bad_password(self):
        for item in [{'name':'bad.exe','data':'YQ=='},{'name':'empty.pdf','data':''},{'name':'bad.pdf','data':'not-base64'}]:
            with self.assertRaises(ValueError):self.call('markdown',item)
        with self.assertRaisesRegex(ValueError,'8–256'):
            self.call('protect',options={'newPassword':'short'})

    @unittest.skipUnless(shutil.which('libreoffice'), 'LibreOffice integration requires container')
    def test_office_conversion(self):
        from docx import Document
        document=Document()
        document.add_paragraph('Office conversion example')
        stream=io.BytesIO()
        document.save(stream)
        item={'name':'example.docx','data':base64.b64encode(stream.getvalue()).decode()}
        data,_,_=self.call('office-to-pdf',item)
        self.assertIn('Office conversion example',PdfReader(io.BytesIO(data)).pages[0].extract_text())

    @unittest.skipUnless(shutil.which('qpdf'), 'qpdf integration requires container')
    def test_repair_rewrites_readable_pdf(self):
        data,_,_=self.call('repair')
        self.assertIn('First document',PdfReader(io.BytesIO(data)).pages[0].extract_text())

    @unittest.skipUnless(shutil.which('gs'), 'Ghostscript integration requires container')
    def test_pdfa_contains_output_intent(self):
        data,_,_=self.call('pdfa')
        doc=PdfReader(io.BytesIO(data))
        self.assertIn('/OutputIntents',doc.trailer['/Root'])
        self.assertIn('/Metadata',doc.trailer['/Root'])

    @unittest.skipUnless(shutil.which('pdftoppm'), 'Poppler integration requires container')
    def test_redaction_discards_text_and_hidden_content(self):
        data,_,_=self.call('redact',options={'box':[0,0,100,100]})
        doc=PdfReader(io.BytesIO(data))
        self.assertEqual(len(doc.pages),1)
        self.assertEqual(doc.pages[0].extract_text(),'')
        self.assertNotIn(b'First document',data)

if __name__=='__main__':unittest.main()
