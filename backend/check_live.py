"""Run explicitly authorized synthetic-file checks against a deployed service."""
import base64
import io
import json
import sys
import urllib.request
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

endpoint = sys.argv[1].rstrip('/')
def item(name, data):
    return {'name':name,'data':base64.b64encode(data).decode()}
def call(task, file, options=None, second=None):
    data=json.dumps({'tool':task,'file':file,'options':options or {},'second':second}).encode()
    request=urllib.request.Request(endpoint+'/process',data=data,headers={'Content-Type':'application/json','Origin':'https://calculatorsallinone.com'})
    with urllib.request.urlopen(request,timeout=120) as response:
        return response.read()
stream=io.BytesIO()
pdf=canvas.Canvas(stream)
pdf.drawString(50,750,'Synthetic PDF test')
pdf.showPage()
pdf.save()
source=item('test.pdf',stream.getvalue())
encrypted=call('protect',source,{'newPassword':'synthetic-check-123'})
doc=PdfReader(io.BytesIO(encrypted))
assert doc.is_encrypted and not doc.decrypt('wrong')
assert doc.decrypt('synthetic-check-123')
unlocked=call('unlock',item('locked.pdf',encrypted),{'password':'synthetic-check-123'})
assert not PdfReader(io.BytesIO(unlocked)).is_encrypted
print('PASS hosted protect and unlock',flush=True)
for task in ['repair','pdfa','redact']:
    result=call(task,source,{'box':[0,0,100,100]} if task=='redact' else {})
    document=PdfReader(io.BytesIO(result))
    assert len(document.pages)==1
    if task=='redact':assert document.pages[0].extract_text()==''
    if task=='pdfa':assert '/OutputIntents' in document.trailer['/Root']
    print('PASS hosted '+task,flush=True)
for task in ['pdf-to-word','pdf-to-excel','pdf-to-powerpoint','markdown','compare']:
    result=call(task,source,second=source if task=='compare' else None)
    assert len(result)>20
    print('PASS hosted '+task,flush=True)
office=[]
doc=Document();doc.add_paragraph('Synthetic Word conversion');stream=io.BytesIO();doc.save(stream);office.append(item('test.docx',stream.getvalue()))
book=Workbook();book.active['A1']='Synthetic Excel conversion';stream=io.BytesIO();book.save(stream);office.append(item('test.xlsx',stream.getvalue()))
deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(1)).text='Synthetic PowerPoint conversion';stream=io.BytesIO();deck.save(stream);office.append(item('test.pptx',stream.getvalue()))
for file in office:
    result=call('office-to-pdf',file)
    assert 'Synthetic' in '\n'.join(p.extract_text() or '' for p in PdfReader(io.BytesIO(result)).pages)
    print('PASS hosted '+file['name'].split('.')[-1]+' to PDF',flush=True)
